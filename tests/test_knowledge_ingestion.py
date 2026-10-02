"""P4b: knowledge base ingestion and the knowledge graph metadata layer.

The controlled document is read into structured Postgres rows (SQLite shim
here), repeatably, versioned by release, with governance on top. Live
reasoning is untouched.

Two inputs:
  the real controlled document, when present (it is not in git: *.docx is
  ignored), for the counts and the spec examples;
  a synthetic document with the same layout, built here, so parsing,
  idempotency and versioning are tested on every run.

Run:  python -m unittest tests.test_knowledge_ingestion -v
"""
from __future__ import annotations

import json
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from fastapi.testclient import TestClient

from pcn_appeal import api
from pcn_appeal.engines.knowledge_matcher import SUPPORTED, KnowledgeMatcher
from pcn_appeal.kg.graph import KnowledgeGraph
from pcn_appeal.knowledge_ingestion import drift, parse
from pcn_appeal.knowledge_ingestion.extract import Vocabulary, extract
from pcn_appeal.knowledge_ingestion.graph import build_graph
from pcn_appeal.knowledge_ingestion.store import DEFAULT_DOCUMENT
from pcn_appeal.manifest import kb_digest
from pcn_appeal.models import CaseFile, EvidenceItem, Fact, FactSource, FactStatus, SourceKind

import sqlite_store

KG = KnowledgeGraph()
REAL = DEFAULT_DOCUMENT
HAVE_REAL = REAL.is_file()

KNOWLEDGE_TABLES = ("knowledge_release", "knowledge_modules", "knowledge_rules",
                    "knowledge_required_facts", "knowledge_evidence_requirements",
                    "knowledge_restrictions", "graph_nodes", "graph_edges",
                    "knowledge_release_items", "knowledge_changes")


# ------------------------------------------------------------------ synthetic document
def _module(d, heading, fields):
    d.add_heading(heading, level=2)
    t = d.add_table(rows=1 + len(fields), cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Field", "Knowledge"
    for i, (k, v) in enumerate(fields, 1):
        t.cell(i, 0).text, t.cell(i, 1).text = k, v


def make_docx(path: Path, *, pay_core="Lead with payment and require reconciliation.",
              extra_module=True, version="1") -> Path:
    import docx
    d = docx.Document()
    d.add_paragraph("TEST PARKING KNOWLEDGE BASE")
    d.add_paragraph(f"VERSION {version} - TEST")
    d.add_heading("5. Payment, Keying and Systems", level=1)
    _module(d, "KB-PAY-01 - Payment made", [
        ("USE WHEN", "Evidence or confirmed facts show a tariff/payment was made."),
        ("CORE PROPOSITION", pay_core),
        ("AI MUST CHECK", "Payment method, time, amount, location, VRM entered."),
        ("EVIDENCE", "Receipt, app confirmation, bank transaction."),
    ])
    d.add_heading("6. ANPR and Evidence Integrity", level=1)
    _module(d, "KB-ANPR-01 - Multiple visits", [
        ("USE WHEN", "Camera records may pair two separate visits."),
        ("CORE PROPOSITION", "The operator must show one continuous stay."),
        ("AI MUST CHECK", "multiple visits; entry time; exit time."),
        ("DO NOT", "Do not raise a generic calibration point. State the specific gap."),
    ])
    if extra_module:
        d.add_heading("99. A New Section", level=1)
        _module(d, "KB-ZZZ-01 - A module nobody coded for", [
            ("USE WHEN", "Something new applies."),
            ("CORE PROPOSITION", "A new proposition."),
            ("AI MUST CHECK", "Event date; relevant land."),
            ("REVIEW NOTE", "A field the parser has never seen."),
        ])
    d.add_heading("17. Mandatory Validator Rules", level=1)
    t = d.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("Validator", "Block release if..."),
                                ("VAL-DRIVER", "Draft identifies the driver."),
                                ("VAL-FACT", "A fact cannot be traced.")]):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    d.add_heading("APPENDIX A - Approved Drafting Building Blocks", level=1)
    d.add_heading("PP-TEST-001 - Synthetic block", level=2)
    d.add_paragraph("THIS DRAFTING PARAGRAPH MUST NEVER BE STORED AS KNOWLEDGE.")
    d.add_heading("APPENDIX C - Release Checklist", level=1)
    d.add_paragraph("Evidence said to be enclosed actually exists.", style="List Bullet")
    d.save(str(path))
    return path


class _Tmp(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="kbingest-"))
        self.addCleanup(shutil.rmtree, self.tmp, True)


def dump_all(db) -> str:
    out = []
    for t in KNOWLEDGE_TABLES:
        out += [repr(r) for r in db.execute(f"SELECT * FROM {t}").fetchall()]
    return "\n".join(out)


def counts(db) -> dict:
    return {t: sqlite_store.count(db, t) for t in KNOWLEDGE_TABLES}


# ================================================================== parsing (no database)
class Parsing(_Tmp):

    def test_structure_is_read_not_assumed(self):
        doc = parse(make_docx(self.tmp / "kb.docx"))
        self.assertEqual(doc.version, "V1")
        self.assertEqual([m.module_id for m in doc.modules], ["KB-PAY-01", "KB-ANPR-01", "KB-ZZZ-01"])
        new = doc.module("KB-ZZZ-01")
        self.assertEqual((new.name, new.family, new.section),
                         ("A module nobody coded for", "ZZZ", "99. A New Section"))
        self.assertIn("REVIEW NOTE", new.fields)          # an unknown field is kept
        self.assertEqual(doc.block_ids, {"PP-TEST-001": "Synthetic block"})
        self.assertEqual([r.rows[0][0] for r in doc.rule_sets], ["VAL-DRIVER"])
        self.assertEqual(len(doc.checklists["APPENDIX C - Release Checklist"]), 1)

    def test_a_module_the_code_has_never_seen_is_extracted_like_any_other(self):
        recs = {r.module_id: r for r in extract(parse(make_docx(self.tmp / "kb.docx")), KG)}
        z = recs["KB-ZZZ-01"]
        self.assertEqual((z.category, z.version, z.status), ("ZZZ", "V1", "ACTIVE"))
        self.assertEqual({f["fact_name"] for f in z.required_facts},
                         {"parking_event_date", "relevant_land"})
        self.assertIn("DOCUMENT_FIELD", {r["rule_type"] for r in z.rules})
        self.assertFalse(z.metadata["compiled"])

    def test_prose_maps_onto_system_vocabulary(self):
        v = Vocabulary.from_kg(KG)
        self.assertEqual(v.facts_in("Payment method"), ["payment_method"])
        self.assertEqual(v.facts_in("event date"), ["parking_event_date"])     # alias
        self.assertEqual(v.facts_in("VRM entered"), ["vrm_entered"])           # longest wins
        self.assertEqual(v.facts_in("operator photos"), [])                    # one word != whole item
        self.assertEqual(v.evidence_kind("bank transaction"), "BANK_STATEMENT")
        self.assertIsNone(v.evidence_kind("signed statement from the moon"))

    def test_no_drafting_paragraph_is_extracted(self):
        doc = parse(make_docx(self.tmp / "kb.docx"))
        recs = extract(doc, KG)
        nodes, edges = build_graph(recs, KG.relations, doc)
        blob = json.dumps([r.as_dict() for r in recs] + [n.as_row() for n in nodes], default=str)
        self.assertNotIn("MUST NEVER BE STORED", blob)


@unittest.skipUnless(HAVE_REAL, "controlled document not present (it is not in git)")
class RealDocumentExtraction(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.doc = parse(REAL)
        cls.recs = {r.module_id: r for r in extract(cls.doc, KG)}
        cls.nodes, cls.edges = build_graph(list(cls.recs.values()), KG.relations, cls.doc)

    def test_every_document_module_is_extracted(self):
        self.assertEqual(len(self.doc.modules), 51)
        self.assertEqual({m.module_id for m in self.doc.modules} - set(self.recs), set())
        # every live module is stored too; the four the document lacks are flagged.
        self.assertEqual(set(self.recs), set(KG.modules) | {m.module_id for m in self.doc.modules})

    def test_spec_2_every_module_has_id_version_category(self):
        for r in self.recs.values():
            with self.subTest(module=r.module_id):
                self.assertTrue(r.module_id and r.version and r.category)
        self.assertEqual(self.recs["KB-POFA-04"].name, "Mandatory notice content")
        self.assertEqual(self.recs["KB-POFA-04"].category, "POFA")
        self.assertEqual(self.recs["KB-POFA-04"].version, "V1")

    def test_no_building_block_text_is_extracted(self):
        texts = [b.text[:60] for b in KG.blocks.values() if len(b.text or "") > 60]
        blob = json.dumps([r.as_dict() for r in self.recs.values()]
                          + [n.as_row() for n in self.nodes], default=str)
        for t in texts:
            self.assertNotIn(t, blob)

    def test_drift_names_the_live_modules_the_document_lacks(self):
        d = drift.report(self.doc, list(self.recs.values()), KG)
        high = {x["module_id"] for x in d if x["severity"] == "HIGH"}
        self.assertEqual(high, set(KG.modules) - {m.module_id for m in self.doc.modules})


# ================================================================== store
class _Store(_Tmp):
    def setUp(self):
        super().setUp()
        self.db = sqlite_store.install(self)
        from pcn_appeal.knowledge_ingestion import store
        self.k = store

    def ingest(self, path, **kw):
        return self.k.ingest(path, created_by=kw.pop("by", "legal.a"),
                             reason=kw.pop("reason", "import"), kg=KG, **kw)


@unittest.skipUnless(HAVE_REAL, "controlled document not present (it is not in git)")
class RealDocumentImport(_Store):

    def setUp(self):
        super().setUp()
        self.result = self.ingest(REAL)

    def test_spec_1_imports_successfully(self):
        self.assertTrue(self.result["created"])
        self.assertEqual(self.result["module_count"], 55)
        self.assertEqual(sqlite_store.count(self.db, "knowledge_modules"), 55)
        self.assertEqual(sqlite_store.count(self.db, "knowledge_release"), 1)

    def test_spec_3_required_facts_are_stored(self):
        m = self.k.get_module("KB-PAY-01")
        req = {(f["fact_name"], f["requirement_type"]) for f in m["required_facts"]}
        self.assertIn(("payment_made", "GATE"), req)
        self.assertIn(("payment_method", "CHECK"), req)
        pofa = {f["fact_name"] for f in self.k.get_module("KB-POFA-01")["required_facts"]}
        self.assertTrue({"driver_status", "notice_route"} <= pofa)

    def test_spec_4_evidence_requirements_are_stored(self):
        ev = {e["evidence_type"]: e["requirement"] for e in self.k.get_module("KB-PAY-01")["evidence"]}
        self.assertTrue({"receipt", "app_confirmation", "bank_transaction"} <= set(ev))
        self.assertEqual(ev["bank_transaction"]["evidence_kind"], "BANK_STATEMENT")

    def test_spec_5_restrictions_are_stored(self):
        r = self.k.get_module("KB-POFA-01")["restrictions"]
        self.assertTrue(any("automatically invalid" in x["content"] and
                            x["restriction_type"] == "PROHIBITED_CLAIM" for x in r))

    def test_spec_6_relationships_are_created(self):
        self.assertGreater(self.result["relationship_count"], 400)
        rels = {(e["source"], e["relationship"]) for e in self.k.get_module("KB-PAY-01")["relationships"]}
        self.assertIn(("FACT:payment_made", "SUPPORTS"), rels)
        anpr = {(e["source"], e["relationship"]) for e in self.k.get_module("KB-ANPR-01")["relationships"]}
        self.assertIn(("EVIDENCE:ATTENDANT_PHOTO", "BLOCKS"), anpr)

    def test_spec_7_duplicate_import_does_not_duplicate(self):
        before = counts(self.db)
        again = self.ingest(REAL)
        self.assertFalse(again["created"])
        self.assertEqual(again["release_id"], self.result["release_id"])
        self.assertEqual(counts(self.db), before)

    def test_spec_9_fact_to_knowledge_matching(self):
        from pcn_appeal.knowledge_ingestion.queries import knowledge_for_facts, modules_for_fact
        rows = [r for r in knowledge_for_facts({"children_present": True})
                if r["relationship"] == "SUPPORTS"]
        bay = [r for r in rows if r["module_id"] == "KB-BAY-02"]
        self.assertTrue(bay, rows)
        self.assertEqual(bay[0]["via"], ["child_occupant_present", "account_contradicts_allegation"])
        self.assertIn("KB-PAY-01", modules_for_fact("payment_made"))
        self.assertEqual(knowledge_for_facts({"children_present": True}),
                         knowledge_for_facts({"children_present": True}))

    def test_no_drafting_paragraph_reaches_any_table(self):
        blob = dump_all(self.db)
        for b in KG.blocks.values():
            if len(b.text or "") > 60:
                self.assertNotIn(b.text[:60], blob, b.block_id)

    def test_a_case_can_name_the_knowledge_release_it_ran_on(self):
        rel = self.k.release_for_digest(kb_digest(KG))
        self.assertEqual(rel["release_id"], self.result["release_id"])
        self.assertEqual(rel["relations_version"], KG.relations.version)


class SyntheticImportAndVersioning(_Store):

    def test_import_requires_a_user_and_a_reason(self):
        path = make_docx(self.tmp / "kb.docx")
        for who, why in (("", "r"), ("a", "")):
            with self.subTest(who=who, why=why), self.assertRaises(self.k.KnowledgeChangeError):
                self.k.ingest(path, created_by=who, reason=why, kg=KG)

    def test_spec_7_twice_is_once(self):
        path = make_docx(self.tmp / "kb.docx")
        first = self.ingest(path)
        before = counts(self.db)
        second = self.ingest(path)
        self.assertEqual((first["created"], second["created"]), (True, False))
        self.assertEqual(counts(self.db), before)

    def test_spec_8_a_changed_document_is_a_new_release_and_can_be_compared(self):
        v1 = self.ingest(make_docx(self.tmp / "v1.docx"))
        v2 = self.ingest(make_docx(self.tmp / "v2.docx", version="2",
                                   pay_core="Lead with payment; reconcile every record.",
                                   extra_module=False))
        self.assertTrue(v2["created"])
        self.assertNotEqual(v1["release_id"], v2["release_id"])
        releases = self.k.list_releases()
        self.assertEqual(releases[0]["parent_release_id"], v1["release_id"])
        self.assertEqual(releases[0]["document_version"], "V2")
        cmp = self.k.compare_releases(v1["release_id"], v2["release_id"])
        self.assertIn("KB-PAY-01", cmp["modules_changed"])
        self.assertEqual(cmp["modules_removed"], ["KB-ZZZ-01"])
        self.assertEqual(self.k.get_module("KB-ZZZ-01")["status"], "RETIRED")
        self.assertTrue(cmp["relationships_removed"])      # ZZZ's DOCX REQUIRES edges
        self.assertEqual(v2["changes"]["modules_retired"], 1)
        rc = self.k.relationship_changes(v2["release_id"])
        self.assertEqual(rc["parent_release_id"], v1["release_id"])
        self.assertEqual(rc["removed"], cmp["relationships_removed"])

    def test_a_changed_compiled_kb_is_a_new_release_too(self):
        path = make_docx(self.tmp / "kb.docx")
        self.ingest(path)
        other = KnowledgeGraph()
        other.modules["KB-PAY-01"].strength += 1
        r = self.k.ingest(path, created_by="a", reason="kb changed", kg=other)
        self.assertTrue(r["created"])

    def test_every_change_is_logged_with_user_time_reason_version(self):
        r = self.ingest(make_docx(self.tmp / "kb.docx"), by="legal.a", reason="first import")
        log = {(c["entity_type"], c["action"]) for c in self.k.list_changes()}
        self.assertIn(("RELEASE", "IMPORT"), log)
        self.assertIn(("MODULE", "CREATE"), log)
        c = self.k.list_changes(r["release_id"])[0]
        self.assertEqual((c["changed_by"], c["reason"], c["version"]), ("legal.a", "first import", "V1"))
        self.assertTrue(c["changed_at"])


class Governance(_Store):

    def setUp(self):
        super().setUp()
        self.path = make_docx(self.tmp / "kb.docx")
        self.ingest(self.path)

    def test_view_and_search(self):
        self.assertEqual({m["module_id"] for m in self.k.list_modules(q="nobody coded")},
                         {"KB-ZZZ-01"})
        self.assertTrue(all(m["category"] == "PAY" for m in self.k.list_modules(category="pay")))
        self.assertIn("KB-PAY-01", {m["module_id"] for m in self.k.list_modules(q="reconciliation")})

    def test_activate_and_disable_need_user_and_reason_and_are_logged(self):
        with self.assertRaises(self.k.KnowledgeChangeError):
            self.k.disable_module("KB-PAY-01", changed_by="", reason="x")
        with self.assertRaises(self.k.KnowledgeChangeError):
            self.k.disable_module("KB-PAY-01", changed_by="legal.a", reason=" ")
        out = self.k.disable_module("KB-PAY-01", changed_by="legal.a", reason="under review")
        self.assertEqual(out["status"], "DISABLED")
        with self.assertRaises(self.k.KnowledgeChangeError):
            self.k.disable_module("KB-PAY-01", changed_by="legal.a", reason="again")
        self.k.activate_module("KB-PAY-01", changed_by="legal.b", reason="reviewed")
        log = self.k.list_changes("KB-PAY-01")
        self.assertEqual([c["action"] for c in log[:2]], ["ACTIVATE", "DISABLE"])
        self.assertEqual((log[0]["changed_by"], log[0]["reason"]), ("legal.b", "reviewed"))

    def test_a_disabled_module_stays_disabled_when_its_content_is_reimported(self):
        self.k.disable_module("KB-PAY-01", changed_by="legal.a", reason="hold")
        self.ingest(make_docx(self.tmp / "v2.docx", pay_core="Changed wording."))
        self.assertEqual(self.k.get_module("KB-PAY-01")["status"], "DISABLED")

    def test_admin_relationships_can_be_made_and_removed(self):
        out = self.k.create_edge("FACT", "payment_attempt_failed", "SUPPORTS", "KNOWLEDGE",
                                 "KB-PAY-01", changed_by="legal.a", reason="link",
                                 confidence=0.5, edge_reason="a failed attempt is a payment story")
        self.assertEqual((out["origin"], out["status"]), ("ADMIN", "REVIEW"))
        with self.assertRaises(self.k.KnowledgeChangeError):
            self.k.create_edge("FACT", "payment_attempt_failed", "SUPPORTS", "KNOWLEDGE",
                               "KB-PAY-01", changed_by="legal.a", reason="dup")
        self.k.remove_edge(out["edge_id"], changed_by="legal.b", reason="not needed")
        self.assertEqual([c["action"] for c in self.k.list_changes(out["edge_id"])],
                         ["REMOVE", "CREATE"])
        # an admin edge survives re-import untouched
        self.ingest(make_docx(self.tmp / "v2.docx", pay_core="Changed."))
        e = [x for x in self.k.list_edges(include_removed=True) if x["edge_id"] == out["edge_id"]]
        self.assertEqual(e[0]["status"], "REMOVED")

    def test_an_ingested_relationship_cannot_be_removed_by_hand(self):
        ingested = next(e for e in self.k.list_edges() if e["origin"] != "ADMIN")
        with self.assertRaises(self.k.KnowledgeChangeError) as cm:
            self.k.remove_edge(ingested["edge_id"], changed_by="a", reason="r")
        self.assertIn("re-import", str(cm.exception))

    def test_unknown_relationship_or_module_is_refused(self):
        with self.assertRaises(self.k.KnowledgeChangeError):
            self.k.create_edge("FACT", "x", "IMPLIES", "KNOWLEDGE", "KB-PAY-01",
                               changed_by="a", reason="r")
        with self.assertRaises(self.k.KnowledgeChangeError):
            self.k.create_edge("FACT", "x", "SUPPORTS", "KNOWLEDGE", "KB-NOPE-99",
                               changed_by="a", reason="r")

    def test_the_store_does_not_change_live_reasoning(self):
        case = CaseFile("C-KI", evidence={"E1": EvidenceItem("E1", "OTHER", "n.pdf", text="x")})
        case.put(Fact("F-alleged_breach", "alleged_breach", "Failed to pay the tariff",
                      FactStatus.CONFIRMED, FactSource(SourceKind.DOCUMENT, "E1#p1")))
        case.put(Fact("F-payment_made", "payment_made", True, FactStatus.ANSWERED,
                      FactSource(SourceKind.ANSWER, "answer:payment_made")))
        before = KnowledgeMatcher(KG).match(case).trace()
        self.k.disable_module("KB-PAY-01", changed_by="a", reason="staged")
        after = KnowledgeMatcher(KG).match(case).trace()
        self.assertEqual(before, after)
        self.assertEqual(KnowledgeMatcher(KG).match(case).candidates["KB-PAY-01"].status, SUPPORTED)


class AdminApi(_Store):

    def setUp(self):
        super().setUp()
        self.ingest(make_docx(self.tmp / "kb.docx"))
        self.client = TestClient(api.app)
        p = mock.patch.dict(os.environ, {"ADMIN_TOKEN": "t0k"})
        p.start()
        self.addCleanup(p.stop)
        self.h = {"X-Admin-Token": "t0k"}

    def test_admin_only(self):
        for method, url, body in (
                ("get", "/admin/knowledge/modules", None),
                ("get", "/admin/knowledge/modules/KB-PAY-01", None),
                ("get", "/admin/knowledge/releases", None),
                ("get", "/admin/knowledge/graph/facts?facts=payment_made", None),
                ("post", "/admin/knowledge/import", {"changed_by": "a", "reason": "r"}),
                ("post", "/admin/knowledge/modules/KB-PAY-01/disable",
                 {"changed_by": "a", "reason": "r"})):
            with self.subTest(url=url):
                r = getattr(self.client, method)(url, **({"json": body} if body else {}))
                self.assertEqual(r.status_code, 401)

    def test_view_disable_compare_through_the_api(self):
        r = self.client.get("/admin/knowledge/modules?q=payment", headers=self.h)
        self.assertIn("KB-PAY-01", {m["module_id"] for m in r.json()["modules"]})
        r = self.client.post("/admin/knowledge/modules/KB-PAY-01/disable", headers=self.h,
                             json={"changed_by": "legal.a", "reason": ""})
        self.assertEqual(r.status_code, 422)
        r = self.client.post("/admin/knowledge/modules/KB-PAY-01/disable", headers=self.h,
                             json={"changed_by": "legal.a", "reason": "hold"})
        self.assertEqual(r.json()["status"], "DISABLED")
        rid = self.client.get("/admin/knowledge/releases", headers=self.h).json()["releases"][0]["release_id"]
        r = self.client.get(f"/admin/knowledge/releases/compare?a={rid}&b={rid}", headers=self.h)
        self.assertEqual(r.json()["modules_changed"], {})
        r = self.client.get("/admin/knowledge/relationship-changes", headers=self.h)
        self.assertEqual(r.status_code, 200, r.text)

    def test_import_only_reads_documents_in_the_data_directory(self):
        r = self.client.post("/admin/knowledge/import", headers=self.h,
                             json={"changed_by": "a", "reason": "r", "document": "../../etc/passwd"})
        self.assertEqual(r.status_code, 422)

    def test_graph_query_endpoint(self):
        r = self.client.get("/admin/knowledge/graph/facts?facts=payment_made", headers=self.h)
        self.assertIn("KB-PAY-01", {x["module_id"] for x in r.json()["knowledge"]})


class WithoutDatabase(unittest.TestCase):
    def test_store_endpoints_say_the_database_is_needed(self):
        with mock.patch.dict(os.environ, {"DATABASE_URL": "", "ADMIN_TOKEN": ""}):
            r = TestClient(api.app).get("/admin/knowledge/modules")
        self.assertEqual(r.status_code, 503)


if __name__ == "__main__":
    unittest.main()
